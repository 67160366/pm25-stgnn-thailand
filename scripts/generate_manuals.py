# NSC 2026 หมวด 14 - ระบบพยากรณ์และวิเคราะห์แหล่งกำเนิด PM2.5 (Explainable STGNN)
# พัฒนาโดย นายรณชัย ขาวสะอาด ม.บูรพา; สนับสนุนโดย สวทช.
# เผยแพร่ตามต้นฉบับ ไม่รับประกันความเสียหาย; ข้อตกลงฉบับเต็ม (ไทย/อังกฤษ) ดู README.md
"""Generate the NSC 2026 install + user manuals as Word documents.

Content is transcribed verbatim from ``docs/INSTALL.md`` and ``docs/USER_GUIDE.md`` (the two
appendix manuals the NSC booklet requires uploaded to SIMS as separate files alongside the final
report). Formatting follows the same booklet rules as ``scripts/generate_report.py``: TH Sarabun
New 16pt, A4, 1-inch margins, an official-style cover (booklet PDF p.44 layout), and footer page
numbers. A lightweight Markdown renderer converts headings, paragraphs, fenced code blocks,
tables, block quotes, and bullet/numbered lists into styled Word blocks.

    UV_NO_SYNC=1 uv run python scripts/generate_manuals.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FONT = "TH Sarabun New"
PROJECT_CODE = "28P14E01196"
TITLE_TH = (
    "ระบบพยากรณ์ฝุ่นละออง PM2.5 และวิเคราะห์แหล่งกำเนิดด้วยโครงข่ายกราฟประสาทเทียม"
    "เชิงปริภูมิ-เวลาแบบอธิบายได้ สำหรับภาคเหนือของประเทศไทย"
)
TITLE_EN = (
    "Explainable Spatio-Temporal Graph Neural Network for PM2.5 Forecasting and "
    "Source Attribution in Northern Thailand"
)

# Manuals to build: (source markdown, cover document-type label, output filename)
MANUALS = [
    ("docs/INSTALL.md", "คู่มือการติดตั้ง (Installation Guide)", "NSC2026_Install_Manual.docx"),
    ("docs/USER_GUIDE.md", "คู่มือการใช้งาน (User Guide)", "NSC2026_User_Manual.docx"),
]

# Emoji / symbol ranges that TH Sarabun New cannot render (strip so no tofu glyphs in a formal doc)
_EMOJI = re.compile("[\U0001f000-\U0001faff☀-➿️]")


# ── Font helpers (pattern from generate_report.py) ──────────────────────────────


def _set_cs_font(run, name: str = FONT) -> None:
    """Force the complex-script font so Thai glyphs render as TH Sarabun New."""
    rpr = run._r.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    existing = rpr.find(qn("w:rFonts"))
    if existing is not None:
        rpr.remove(existing)
    rpr.insert(0, rfonts)


def fmt(run, size: int = 16, bold: bool = False, italic: bool = False, font: str = FONT) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    _set_cs_font(run, font)


def _apply_cs_to_style(style, name: str = FONT) -> None:
    """Set ascii/hAnsi/cs font on a style definition (so Thai renders in headings/Normal)."""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)


# ── Inline markdown (**bold**, *italic*, `code`, [text](url)) → runs ─────────────

_INLINE = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")
_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_md_runs(p, text: str, size: int = 16, italic: bool = False) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        link = _LINK.match(part) if part.startswith("[") else None
        if link:
            label, url = link.group(1), link.group(2)
            shown = label if label == url else f"{label} ({url})"
            r = p.add_run(shown)
            fmt(r, size=size, italic=italic)
        elif part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            fmt(r, size=size, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            fmt(r, size=size, font="Courier New")
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1])
            fmt(r, size=size, italic=True)
        else:
            r = p.add_run(part)
            fmt(r, size=size, italic=italic)


# ── Block helpers ───────────────────────────────────────────────────────────────


def para(
    doc, text="", size=16, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        fmt(r, size=size, bold=bold, italic=italic)
    return p


def body(doc, text, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    add_md_runs(p, text, size=size)
    return p


def h1(doc, text, size=18):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    fmt(r, size=size, bold=True)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def h2(doc, text, size=16):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    fmt(r, size=size, bold=True)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def h3(doc, text, size=16):
    p = doc.add_paragraph(style="Heading 3")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    fmt(r, size=size, bold=True, italic=True)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def bullet(doc, text, size=16, indent_cm=1.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(indent_cm)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("• ")
    fmt(r, size=size)
    add_md_runs(p, text, size=size)
    return p


def numbered(doc, n, text, size=16, indent_cm=0.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(indent_cm)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{n}. ")
    fmt(r, size=size)
    add_md_runs(p, text, size=size)
    return p


def callout(doc, text, size=15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    add_md_runs(p, text, size=size)
    return p


def _shade(p, fill: str = "F2F2F2") -> None:
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def code_block(doc, code_lines, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    _shade(p)
    for i, line in enumerate(code_lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        fmt(r, size=size, font="Courier New")
    return p


def _cell(cell, text, size=14, bold=False, center=False):
    cp = cell.paragraphs[0]
    if center:
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_md_runs(cp, text, size=size)
    if bold:
        for r in cp.runs:
            r.font.bold = True


def table(doc, headers, rows, size=14):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for j, hdr in enumerate(headers):
        _cell(t.rows[0].cells[j], hdr, size=size, bold=True, center=True)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            if j < len(headers):
                _cell(t.rows[i].cells[j], val, size=size)
    para(doc, "", size=6)
    return t


# ── Footer page number (booklet rule) ───────────────────────────────────────────


def add_page_number(section):
    section.different_first_page_header_footer = True
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)
    fmt(run, size=14)


# ── Cover (booklet final-report template layout, PDF p.44) ───────────────────────


def build_cover(doc, doc_type: str) -> None:
    c = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, f"รหัสโครงการ  {PROJECT_CODE}", size=16, align=c, space_after=18)
    para(doc, TITLE_TH, size=20, bold=True, align=c, space_after=2)
    para(doc, f"({TITLE_EN})", size=15, italic=True, align=c, space_after=14)
    para(
        doc,
        "หมวด 14: โปรแกรมเพื่องานการพัฒนาด้านวิทยาศาสตร์และเทคโนโลยี",
        size=16,
        bold=True,
        align=c,
        space_after=2,
    )
    para(doc, "(ระดับนิสิต นักศึกษา)", size=16, align=c, space_after=18)
    para(doc, doc_type, size=20, bold=True, align=c, space_after=4)
    para(doc, "(ภาคผนวกของรายงานฉบับสมบูรณ์)", size=15, italic=True, align=c, space_after=14)
    para(doc, "เสนอต่อ", size=16, align=c, space_after=2)
    para(
        doc,
        "สำนักงานพัฒนาวิทยาศาสตร์และเทคโนโลยีแห่งชาติ",
        size=16,
        bold=True,
        align=c,
        space_after=2,
    )
    para(doc, "กระทรวงการอุดมศึกษา วิทยาศาสตร์ วิจัยและนวัตกรรม", size=16, align=c, space_after=14)
    para(doc, "ได้รับทุนอุดหนุนโครงการวิจัย พัฒนาและวิศวกรรม", size=16, align=c, space_after=2)
    para(
        doc,
        "โครงการแข่งขันพัฒนาโปรแกรมคอมพิวเตอร์แห่งประเทศไทย ครั้งที่ 28",
        size=16,
        bold=True,
        align=c,
        space_after=2,
    )
    para(doc, "ประจำปีงบประมาณ 2569", size=16, align=c, space_after=18)
    para(doc, "โดย", size=16, align=c, space_after=2)
    para(doc, "นายรณชัย ขาวสะอาด (หัวหน้าโครงการ)", size=16, bold=True, align=c, space_after=10)
    para(doc, "อาจารย์ที่ปรึกษาโครงการ: ดร.วัชรพงศ์ อยู่ขวัญ", size=16, align=c, space_after=2)
    para(
        doc,
        "คณะวิทยาการสารสนเทศ มหาวิทยาลัยบูรพา  ·  จังหวัดชลบุรี",
        size=16,
        align=c,
        space_after=2,
    )
    doc.add_page_break()


# ── Markdown → Word renderer ─────────────────────────────────────────────────────

_BULLET = re.compile(r"^(\s*)[-*]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
_HR = re.compile(r"^([-*_])\1{2,}$")
_SEP_CELL = re.compile(r"^:?-{2,}:?$")


def _clean(text: str) -> str:
    return _EMOJI.sub("", text).rstrip()


def _render_table(doc, rows: list[str]) -> None:
    parsed = []
    for raw in rows:
        cells = [c.strip() for c in raw.strip().strip("|").split("|")]
        if all(_SEP_CELL.match(c) for c in cells):
            continue  # header/body separator row
        parsed.append(cells)
    if not parsed:
        return
    width = max(len(r) for r in parsed)
    parsed = [r + [""] * (width - len(r)) for r in parsed]
    table(doc, parsed[0], parsed[1:])


def render_markdown(doc, md: str) -> None:
    lines = md.split("\n")
    para_buf: list[str] = []
    quote_buf: list[str] = []
    table_buf: list[str] = []

    def flush_para():
        if para_buf:
            body(doc, " ".join(para_buf))
            para_buf.clear()

    def flush_quote():
        if quote_buf:
            callout(doc, " ".join(quote_buf))
            quote_buf.clear()

    def flush_table():
        if table_buf:
            _render_table(doc, table_buf)
            table_buf.clear()

    def flush_all():
        flush_para()
        flush_quote()
        flush_table()

    i = 0
    while i < len(lines):
        raw = _clean(lines[i])
        stripped = raw.strip()

        # Fenced code block
        if stripped.startswith("```"):
            flush_all()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(_clean(lines[i]))
                i += 1
            code_block(doc, code)
            i += 1
            continue

        # Table rows accumulate; anything else flushes the pending table
        if stripped.startswith("|"):
            flush_para()
            flush_quote()
            table_buf.append(raw)
            i += 1
            continue
        flush_table()

        if not stripped:
            flush_para()
            flush_quote()
        elif stripped.startswith("#"):
            flush_all()
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip().replace("`", "").replace("**", "")
            if level <= 1:
                pass  # document title already on the cover
            elif level == 2:
                h1(doc, text)
            elif level == 3:
                h2(doc, text)
            else:
                h3(doc, text)
        elif stripped.startswith(">"):
            flush_para()
            quote_buf.append(stripped.lstrip(">").strip())
        elif _HR.match(stripped):
            flush_all()
        elif _BULLET.match(raw):
            flush_para()
            flush_quote()
            m = _BULLET.match(raw)
            indent = 1.5 if m.group(1) else 1.0
            bullet(doc, m.group(2), indent_cm=indent)
        elif _NUMBERED.match(raw):
            flush_para()
            flush_quote()
            m = _NUMBERED.match(raw)
            numbered(doc, m.group(2), m.group(3))
        else:
            flush_quote()
            para_buf.append(stripped)
        i += 1

    flush_all()


# ════════════════════════════════════════════════════════════════════════════════
# BUILD
# ════════════════════════════════════════════════════════════════════════════════


def build_manual(md_path: Path, doc_type: str, out_path: Path) -> None:
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(16)
    _apply_cs_to_style(normal)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        try:
            _apply_cs_to_style(doc.styles[style_name])
        except KeyError:
            pass

    add_page_number(sec)
    build_cover(doc, doc_type)
    render_markdown(doc, md_path.read_text(encoding="utf-8"))
    doc.save(str(out_path))


def main() -> None:
    for md_rel, doc_type, out_name in MANUALS:
        md_path = ROOT / md_rel
        out_path = ROOT / "outputs" / out_name
        build_manual(md_path, doc_type, out_path)
        print(f"wrote {out_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
