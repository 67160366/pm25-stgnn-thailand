# NSC 2026 หมวด 14 - ระบบพยากรณ์และวิเคราะห์แหล่งกำเนิด PM2.5 (Explainable STGNN)
# พัฒนาโดย นายรณชัย ขาวสะอาด ม.บูรพา; สนับสนุนโดย สวทช.
# เผยแพร่ตามต้นฉบับ ไม่รับประกันความเสียหาย; ข้อตกลงฉบับเต็ม (ไทย/อังกฤษ) ดู README.md
"""Generate the AI Camp pitch deck (9 slides) as an editable .pptx.

Mirrors the CORRECTED `outputs/pitch/presenter_script_th.md`: leads with XAI / source
attribution and engineering honesty, uses the corrected val-2025 numbers (NOT the discredited
denorm-bug numbers 10.21 / +5.1% / +10.6%), and adds the held-out Mae Hong Son transboundary
case. Every number here traces to a verified source JSON in `outputs/`.

Run: ``uv run python scripts/generate_pitch_deck.py`` -> writes outputs/pitch/pitch_deck_aicamp.pptx
The submitted ``docs/project_pitch.pdf`` is left untouched (immutable history).
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.slide import Slide
from pptx.text.text import TextFrame, _Run
from pptx.util import Inches, Pt

# ── palette ────────────────────────────────────────────────────────────────────
ACCENT = RGBColor(0x15, 0x61, 0x6D)  # deep teal (title bars / title slide)
ACCENT2 = RGBColor(0xE8, 0x74, 0x3B)  # orange (kicker / emphasis)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Leelawadee"  # Thai-capable system font (falls back to Tahoma if absent)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

OUTPUT = Path(__file__).resolve().parents[1] / "outputs" / "pitch" / "pitch_deck_aicamp.pptx"


def _style_run(run: _Run, *, size: int, bold: bool = False, color: RGBColor = DARK) -> None:
    """Apply size/bold/color and force latin+ea+cs typefaces so Thai renders correctly."""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        for existing in rpr.findall(qn(tag)):
            rpr.remove(existing)
        el = rpr.makeelement(qn(tag), {"typeface": FONT})
        rpr.append(el)


def _add_box(
    slide: Slide,
    left: float,
    top: float,
    width: float,
    height: float,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> TextFrame:
    """Add a text frame (inches) and return it, word-wrap on, given vertical anchor."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def _bar(
    slide: Slide, left: float, top: float, width: float, height: float, color: RGBColor
) -> None:
    """Add a filled rectangle (no outline) as a colored band."""
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _footer(slide: Slide, page: int) -> None:
    tf = _add_box(slide, 0.5, 7.0, 12.3, 0.4)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "ระบบพยากรณ์ PM2.5 + วิเคราะห์แหล่งกำเนิด (Explainable STGNN)  |  NSC 2026 หมวด 14"
    _style_run(r, size=10, color=MUTED)
    pn = _add_box(slide, 12.4, 7.0, 0.6, 0.4)
    pp = pn.paragraphs[0]
    pp.alignment = PP_ALIGN.RIGHT
    rr = pp.add_run()
    rr.text = str(page)
    _style_run(rr, size=10, color=MUTED)


def add_title_slide(prs: Presentation) -> None:
    """Slide 1: full-accent title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bar(slide, 0, 0, 13.333, 7.5, ACCENT)
    tf = _add_box(slide, 1.0, 2.2, 11.3, 2.6, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "ระบบพยากรณ์ PM2.5 และวิเคราะห์แหล่งกำเนิด\nด้วย Explainable STGNN"
    _style_run(r, size=38, bold=True, color=WHITE)
    sub = tf.add_paragraph()
    sub.space_before = Pt(14)
    rs = sub.add_run()
    rs.text = "สำหรับ 9 จังหวัดภาคเหนือ — พยากรณ์ล่วงหน้า + ระบุแหล่งกำเนิดที่อธิบายได้"
    _style_run(rs, size=20, color=WHITE)
    foot = _add_box(slide, 1.0, 6.3, 11.3, 0.6)
    pf = foot.paragraphs[0]
    rf = pf.add_run()
    rf.text = "NSC 2026 — หมวด 14  |  นำเสนอ 2–3 ก.ค. 2569  |  [ชื่อทีม / ผู้นำเสนอ]"
    _style_run(rf, size=14, color=WHITE)


def add_content_slide(
    prs: Presentation, page: int, kicker: str, title: str, bullets: list[tuple[str, int, bool]]
) -> None:
    """Add a content slide: top accent bar with title + kicker, then bulleted body."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bar(slide, 0, 0, 13.333, 1.25, ACCENT)
    # kicker (small, above/within bar) + title
    head = _add_box(slide, 0.5, 0.12, 12.3, 1.05, anchor=MSO_ANCHOR.MIDDLE)
    pk = head.paragraphs[0]
    rk = pk.add_run()
    rk.text = kicker
    _style_run(rk, size=13, bold=True, color=RGBColor(0xCF, 0xE6, 0xEA))
    pt = head.add_paragraph()
    rt = pt.add_run()
    rt.text = title
    _style_run(rt, size=27, bold=True, color=WHITE)
    # body bullets
    body = _add_box(slide, 0.7, 1.6, 12.0, 5.2)
    first = True
    for text, level, bold in bullets:
        p = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(10)
        glyph = "•  " if level == 0 else "–  "
        r = p.add_run()
        r.text = glyph + text
        _style_run(r, size=18 if level == 0 else 15, bold=bold, color=DARK)
    _footer(slide, page)


def build() -> None:
    """Assemble all 9 slides and save the deck."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_title_slide(prs)

    add_content_slide(
        prs,
        2,
        "ปัญหา & ผลกระทบ (Social)",
        "วิกฤตหมอกควัน PM2.5 ภาคเหนือ",
        [
            ("ทุกฤดูแล้ง ภาคเหนือเผชิญวิกฤตหมอกควัน PM2.5", 0, False),
            (
                "มี.ค. 2024 เชียงใหม่วัดได้ 141–144 µg/m³ — สูงกว่าค่าแนะนำ WHO เกือบ 10 เท่า",
                0,
                True,
            ),
            ("กระทบสุขภาพประชาชนหลายล้านคน + การท่องเที่ยว + เศรษฐกิจ", 0, False),
            (
                "ปัญหาหลัก: ไม่รู้ล่วงหน้า และไม่รู้ว่าฝุ่นมาจากไหน → ตัดสินใจเชิงนโยบายยาก",
                0,
                False,
            ),
        ],
    )

    add_content_slide(
        prs,
        3,
        "ช่องว่างของเครื่องมือปัจจุบัน",
        "ตั้งรับ — ไม่พยากรณ์ ไม่บอกแหล่งที่มา",
        [
            ("เว็บ AQI / การแจ้งเตือน คพ. = ระบบตั้งรับ บอกแค่ค่าปัจจุบัน", 0, False),
            ("ไม่พยากรณ์ล่วงหน้า และไม่ระบุสัดส่วนแหล่งกำเนิดไฟ", 0, False),
            ("เราสร้างระบบที่ทำทั้งสองอย่าง: พยากรณ์ + อธิบายแหล่งที่มา", 0, True),
        ],
    )

    add_content_slide(
        prs,
        4,
        "3 จุดใหม่ (Creativity)",
        "สามนวัตกรรมหลัก",
        [
            ("Wind-aware dynamic graph — เส้นเชื่อมสถานีปรับตามทิศลม ERA5 รายชั่วโมง", 0, False),
            (
                "FIRMS hotspot เป็น node ในกราฟ — จุดความร้อนดาวเทียม NASA เป็นต้นทางไฟโดยตรง",
                0,
                False,
            ),
            (
                "Graph-based Integrated Gradients — อธิบายว่าการพยากรณ์มาจากแหล่งใด/ประเทศใด",
                0,
                False,
            ),
        ],
    )

    add_content_slide(
        prs,
        5,
        "สถาปัตยกรรม & สาธิต",
        "MTGNN + กราฟ 3 ชนิดเส้นเชื่อม",
        [
            (
                "โมเดล MTGNN: 18 สถานี + โหนด hotspot ผ่าน 3 ชนิดเส้นเชื่อม (พื้นที่ / ลม / ไฟ)",
                0,
                False,
            ),
            ("พยากรณ์ 4 ช่วง: 6, 12, 24, 48 ชั่วโมง", 0, False),
            ("[สาธิต dashboard สด: แผนที่พยากรณ์ 18 สถานี → แท็บ Source Attribution]", 0, True),
            ("ระบบรันได้จริง ไม่ใช่แค่ตัวเลขบนกระดาษ", 0, False),
        ],
    )

    add_content_slide(
        prs,
        6,
        "ผลลัพธ์จริง (Technical)",
        "วัดอย่างเป็นระบบ + ซื่อสัตย์",
        [
            ("วัดบนข้อมูลปี 2025 (3,637 ตัวอย่าง, 18 สถานี) เทียบ baseline persistence", 0, False),
            ("48 ชม. (ช่วงเตือนภัยสำคัญสุด): MTGNN RMSE 12.68 vs persistence 12.99 µg/m³", 0, True),
            (
                "ช่วงสั้น 6–12 ชม.: persistence ยังเหนือกว่า (PM2.5 autocorrelation สูง) — ตรงกับงาน PM2.5-GNN",
                0,
                False,
            ),
            (
                "ระบบ Hybrid: persistence ช่วงสั้น + MTGNN ช่วงยาว → ดีเท่าหรือดีกว่า baseline ทุกช่วง",
                0,
                False,
            ),
            (
                "ความซื่อสัตย์: ตรวจพบและแก้ denorm bug หลังส่ง proposal; ทุกผลลัพธ์ reproduce ได้",
                0,
                False,
            ),
        ],
    )

    add_content_slide(
        prs,
        7,
        "XAI / แหล่งกำเนิด  ★ จุดขายหลัก",
        "ระบุแหล่งกำเนิด — สิ่งที่ baseline ทำไม่ได้",
        [
            ("สิ่งที่ persistence / โมเดลทั่วไปทำไม่ได้ = การระบุแหล่งกำเนิด", 0, False),
            (
                "เชียงใหม่ (กลางเมือง) มี.ค. 2024: ไฟในไทยเป็นหลัก ~100% — FRP ไทยสูงกว่าเมียนมา 128×",
                0,
                False,
            ),
            (
                "ช่อง hotspot เพิ่มค่าพยากรณ์ ~3.5 µg/m³; Integrated Gradients → ความชื้น/อุณหภูมิ (d2m, t2m)",
                1,
                False,
            ),
            (
                "แม่ฮ่องสอน (ชายแดน) — held-out 2025-02-16: ไฟใกล้ชายแดนเป็นของเมียนมา ~37% → โมเดลชี้เมียนมา ~37%",
                0,
                True,
            ),
            (
                "สอดคล้องกับสัดส่วนไฟจริง — ระบบจับฝุ่นข้ามแดนได้ “เมื่อ-และที่-ไหน” มันเกิด; กลางเมืองคือไฟไทยจริง",
                1,
                False,
            ),
        ],
    )

    add_content_slide(
        prs,
        8,
        "AI Governance & ข้อจำกัด",
        "ใช้ AI อย่างรับผิดชอบ + โปร่งใส",
        [
            (
                "ERA5 มี latency → ระบบจริงต้องใช้ NWP forecast; วัดผลแล้ว: edge 48h ทนได้เมื่อ noise < 0.5× ความผันผวนธรรมชาติ (ที่ 0.5× เริ่มแพ้)",
                0,
                False,
            ),
            (
                "Attribution = การประมาณจากโมเดล ตรวจเทียบข้อมูลไฟจริง FIRMS — ไม่ใช่การวัดตรง",
                0,
                False,
            ),
            ("สื่อสารพร้อมความไม่แน่นอน และไม่ควรใช้กล่าวโทษเชิงการทูต", 1, False),
            ("ใช้เฉพาะข้อมูลสาธารณะ ไม่มีข้อมูลส่วนบุคคล วิธีการ reproduce ได้ทั้งหมด", 0, False),
        ],
    )

    add_content_slide(
        prs,
        9,
        "สรุป & ต่อยอด",
        "เตือนภัยล่วงหน้า + ระบุแหล่งกำเนิด",
        [
            (
                "ระบบให้การเตือนภัยหมอกควันล่วงหน้า + ระบุแหล่งกำเนิด สำหรับ 9 จังหวัดภาคเหนือ",
                0,
                True,
            ),
            ("ต่อยอด: เพิ่มสถานี คพ., ต่อ NWP forecast จริง, ขยายไปภูมิภาคอื่น", 0, False),
            ("ขอบคุณครับ / ค่ะ", 0, False),
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"Saved deck: {OUTPUT}  ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
